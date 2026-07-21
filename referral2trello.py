import re
import requests
import os
import shutil
import pickle
from docx import Document
import subprocess
from pypdf import PdfReader


from google_auth_oauthlib.flow import InstalledAppFlow
from googleapiclient.discovery import build
from googleapiclient.http import MediaFileUpload

# ================================
# 🔧 CONFIG
# ================================

INPUT_FOLDER = "/mnt/chromeos/MyFiles/Downloads/referrals"
PROCESSED_FOLDER = os.path.join(INPUT_FOLDER, "processed")

TRELLO_API_KEY = os.getenv("TRELLO_API_KEY")
TRELLO_TOKEN = os.getenv("TRELLO_API_TOKEN")
TRELLO_LIST_ID = os.getenv("TRELLO_CREATE_IA_LIST_ID")
DRIVE_FOLDER_ID = os.getenv("DRIVE_FOLDER_ID")

CREDENTIALS_PATH = "/home/springvolunteer/Airtable2Trello/credentials.json"
SCOPES = ['https://www.googleapis.com/auth/drive.file']

os.makedirs(PROCESSED_FOLDER, exist_ok=True)

# ================================
# 🚫 DUPLICATES
# ================================

SEEN_FILE = os.path.join(INPUT_FOLDER, "seen_records.txt")

def load_seen():
    if not os.path.exists(SEEN_FILE):
        return set()
    with open(SEEN_FILE, "r") as f:
        return set(line.strip() for line in f)

def save_seen(record):
    with open(SEEN_FILE, "a") as f:
        f.write(record + "\n")

# ================================
# 🔐 GOOGLE AUTH
# ================================

def get_drive_service():
    creds = None

    if os.path.exists("token.pickle"):
        with open("token.pickle", "rb") as token:
            creds = pickle.load(token)

    if not creds or not creds.valid:
        flow = InstalledAppFlow.from_client_secrets_file(
            CREDENTIALS_PATH, SCOPES
        )
        creds = flow.run_local_server(port=0)

        with open("token.pickle", "wb") as token:
            pickle.dump(creds, token)

    return build("drive", "v3", credentials=creds)


def extract_pdf_fields(file_path):
    reader = PdfReader(file_path)

    text = ""

    for page in reader.pages:
        page_text = page.extract_text()

        if page_text:
            text += page_text + "\n"

    data = {"notes": text}

    phone_match = re.search(
        r"(\+44\s?\d[\d\s]+|0\d[\d\s]{8,})",
        text
    )

    if phone_match:
        data["phone"] = phone_match.group(1)

    email_match = re.search(
        r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}",
        text
    )

    if email_match:
        data["email"] = email_match.group(0)

    # crude name extraction
    lines = [l.strip() for l in text.splitlines() if l.strip()]

    if lines:
        data["name"] = lines[0]
    else:
        data["name"] = "Unknown Client"

    return data


def extract_doc_fields(file_path):

    subprocess.run(
        [
            "libreoffice",
            "--headless",
            "--convert-to",
            "docx",
            file_path,
            "--outdir",
            os.path.dirname(file_path)
        ],
        check=True
    )

    converted = os.path.splitext(file_path)[0] + ".docx"

    return extract_docx_fields(converted)


def extract_pdf_text(file_path):
    text = ""

    reader = PdfReader(file_path)

    for page in reader.pages:
        page_text = page.extract_text()

        if page_text:
            text += page_text + "\n"

    return text

def extract_fields(file_path):

    ext = os.path.splitext(file_path)[1].lower()

    if ext == ".docx":
        return extract_docx_fields(file_path)

    elif ext == ".pdf":
        return extract_pdf_fields(file_path)

    elif ext == ".doc":
        return extract_doc_fields(file_path)

    raise ValueError(f"Unsupported file type: {ext}")


# ================================
# 🧠 FIELD EXTRACTION (FORM AWARE)
# ================================

def extract_docx_fields(file_path):
    doc = Document(file_path)

    data = {"notes": ""}

    # ================================
    # READ TABLES
    # ================================

    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]

            if len(cells) < 2:
                continue

            key = cells[0].lower()
            val = cells[1].strip()

            if not val:
                continue

            if any(x in key for x in [
                "first name",
                "forename",
                "client first name"
            ]):
                data["first_name"] = val

            elif any(x in key for x in [
                "surname",
                "last name",
                "family name"
            ]):
                data["surname"] = val

            elif any(x in key for x in [
                "telephone",
                "phone",
                "mobile",
                "contact",
                "contact number",
                "phone number"
            ]):
                data["phone"] = val

            elif "whatsapp" in key:
                data["whatsapp"] = val

            elif "email" in key:
                data["email"] = val

            elif "address" in key:
                data["address"] = val

            elif "date of birth" in key or "dob" in key:
                data["dob"] = val

            elif any(x in key for x in [
                "form completed by",
                "completed by",
                "referrer"
            ]):
                data["referrer"] = val

            else:
                data["notes"] += val + "\n"

    # ================================
    # READ PARAGRAPHS
    # ================================

    for para in doc.paragraphs:
        text = para.text.strip()

        if text:
            data["notes"] += text + "\n"

    # ================================
    # FALLBACK PHONE SEARCH
    # ================================

    full_text = "\n".join(
        p.text for p in doc.paragraphs
    )

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                full_text += "\n" + cell.text

    if not data.get("phone"):
        phone_match = re.search(
            r"(\+44\s?\d[\d\s]+|0\d[\d\s]{8,})",
            full_text
        )

        if phone_match:
            data["phone"] = phone_match.group(1)

    if not data.get("email"):
        email_match = re.search(
            r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}",
            full_text
        )

        if email_match:
            data["email"] = email_match.group(0)

    # ================================
    # NAME
    # ================================

    first = data.get("first_name", "")
    surname = data.get("surname", "")

    data["name"] = f"{first} {surname}".strip()

    if not data["name"]:
        data["name"] = "Unknown Client"

    return data


    # =================================
    # READ PARAGRAPHS TOO
    # =================================

    for para in doc.paragraphs:
        text = para.text.strip()

        if not text:
            continue

        lower = text.lower()

        if (
            "telephone" in lower
            or "phone" in lower
            or "mobile" in lower
        ):
            if "phone" not in data:
                phone_match = re.search(
                    r"(\+44\s?\d[\d\s]+|0\d[\d\s]{8,})",
                    text
                )

                if phone_match:
                    data["phone"] = phone_match.group(1)

        elif "email" in lower:
            if "email" not in data:
                email_match = re.search(
                    r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}",
                    text
                )

                if email_match:
                    data["email"] = email_match.group(0)

        else:
            data["notes"] += text + "\n"

    # =================================
    # FALLBACK FULL-TEXT SEARCH
    # =================================

    full_text = "\n".join(p.text for p in doc.paragraphs)

    if not data.get("phone"):
        phone_match = re.search(
            r"(\+44\s?\d[\d\s]+|0\d[\d\s]{8,})",
            full_text
        )

        if phone_match:
            data["phone"] = phone_match.group(1)

    if not data.get("email"):
        email_match = re.search(
            r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}",
            full_text
        )

        if email_match:
            data["email"] = email_match.group(0)

    # =================================
    # COMBINE NAME
    # =================================

    first = data.get("first_name", "")
    surname = data.get("surname", "")

    data["name"] = f"{first} {surname}".strip()

    if not data["name"]:
        data["name"] = (
            data.get("first_name")
            or data.get("surname")
            or "Unknown Client"
        )

    return data


# ================================
# 🧹 HELPERS
# ================================

def get_field(data, key, default="unknown"):
    val = data.get(key, "").strip()
    return val if val else default

def clean_text(text):
    return re.sub(r"[^A-Za-z0-9]", "_", text)

def clean_phone(phone):
    digits = re.sub(r"\D", "", phone)
    if digits.startswith("0"):
        digits = "44" + digits[1:]
    return digits or "no_phone"

# ================================
# 📄 FILENAME
# ================================

def generate_filename(data):
    name = clean_text(get_field(data, "name"))
    phone = clean_phone(get_field(data, "phone"))
    return f"{name}_{phone}.docx"[:120]

def ensure_unique(path):
    base, ext = os.path.splitext(path)
    i = 2
    while os.path.exists(path):
        path = f"{base}_{i}{ext}"
        i += 1
    return path

# ================================
# 🚫 DUP CHECK
# ================================

def is_duplicate(data, seen):
    key = f"{clean_text(get_field(data,'name'))}_{clean_phone(get_field(data,'phone'))}"
    if key in seen:
        print(f"🚫 Duplicate: {key}")
        return True, key
    return False, key

# ================================
# ☁️ UPLOAD TO DRIVE
# ================================

def upload_to_drive(file_path):
    print("☁️ Uploading to Drive...")

    service = get_drive_service()

    file_metadata = {
        "name": os.path.basename(file_path),
        "parents": [DRIVE_FOLDER_ID],
    }

    media = MediaFileUpload(file_path)

    file = service.files().create(
        body=file_metadata,
        media_body=media,
        fields="id"
    ).execute()

    link = f"https://drive.google.com/file/d/{file['id']}/view"
    print(f"✅ Uploaded: {link}")
    return link

# ================================
# 🚀 CREATE TRELLO CARD
# ================================

def create_card(data, link):
    name = get_field(data, "name")
    contact_phone = get_field(data, "phone", "")
    whatsapp_phone = get_field(data, "whatsapp", "") or contact_phone

    whatsapp_number = clean_phone(whatsapp_phone)

    # clean referrer formatting
    referrer_raw = data.get("referrer", "").strip()
    referrer = " ".join(referrer_raw.split()) if referrer_raw else "unknown"

    desc = f"""
Name: {name}
Phone: {contact_phone}
Email: {get_field(data,'email')}
DOB: {get_field(data,'dob')}
Address: {get_field(data,'address')}

Referrer:
{referrer}

Notes:
{get_field(data,'notes')}

Other (please explain):
{get_field(data,'notes')}

📱 WhatsApp:
https://wa.me/{whatsapp_number}

💬 WhatsApp Web:
https://web.whatsapp.com/send?phone={whatsapp_number}

📎 File:
{link}
"""

    response = requests.post(
        "https://api.trello.com/1/cards",
        params={
            "key": TRELLO_API_KEY,
            "token": TRELLO_TOKEN,
        },
        data={
            "idList": TRELLO_LIST_ID,
            "name": f"{name} ({contact_phone})",
            "desc": desc[:15000],   # optional safeguard
        }
    )

    print(f"Description length: {len(desc)}")

    if response.status_code == 200:
        print("✅ Trello card created")
    else:
        print("❌ Trello error", response.text)

# ================================
# 🧠 PROCESS FILE
# ================================

def process_file(file_path, seen):
    print(f"\n📄 Processing: {file_path}")

    data = extract_fields(file_path)


    print("\n=== EXTRACTED DATA ===")
    for k, v in data.items():
        print(f"{k}: {v}")
    print("======================\n")


    duplicate, key = is_duplicate(data, seen)
    if duplicate:
        return

    filename = generate_filename(data)
    new_path = ensure_unique(os.path.join(INPUT_FOLDER, filename))

    os.rename(file_path, new_path)
    print(f"✏️ Renamed: {filename}")

    link = upload_to_drive(new_path)
    
    print("\n=== EXTRACTED ===")
    for k, v in data.items():
        print(f"{k}: {str(v)[:200]}")
    print("=================\n")


    create_card(data, link)

    save_seen(key)
    seen.add(key)

    shutil.move(new_path, os.path.join(PROCESSED_FOLDER, os.path.basename(new_path)))
    print("📦 Moved to processed")

# ================================
# ▶️ MAIN
# ================================

def main():
    print("📂 Processing folder...")
    seen = load_seen()

    for f in os.listdir(INPUT_FOLDER):

        if not f.lower().endswith((".docx", ".doc", ".pdf")):
            continue

        try:
            process_file(
                os.path.join(INPUT_FOLDER, f),
                seen
            )
        except Exception as e:
            print(f"❌ Failed {f}: {e}")



# ================================

if __name__ == "__main__":
    main()
